from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.urls import reverse
from django.contrib.auth.decorators import login_required
from django.http import HttpResponseForbidden, JsonResponse, HttpResponse, Http404, FileResponse
from django.db.models import Max
from django.contrib.auth.forms import UserCreationForm, AuthenticationForm
from django.contrib.auth import login, authenticate, logout
from django.core.cache import cache
import random
import os
import mimetypes
from functools import wraps
try:
    from pypdf import PdfReader as _PdfReader
    _PDF_PKG = 'pypdf'
except Exception:
    try:
        from PyPDF2 import PdfReader as _PdfReader
        _PDF_PKG = 'PyPDF2'
    except Exception:
        _PdfReader = None
        _PDF_PKG = None
from docx import Document
import json

from django.contrib.auth.models import User
from django.conf import settings
from django.utils import timezone
from questions.models import Question, UserAnswer, AIDocument, AIChat, UserAIUsage, InterviewSession, InterviewAnswer, K8sProgress, LangChainProgress
from questions.ai_similarity import calculate_similarity, preprocess_text, fallback_calculate_similarity, generate_ai_content


def extract_text_from_file(file_path, file_type):
    """从不同类型的文件中提取文本"""
    text = ""
    
    try:
        if file_type == 'txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
        elif file_type == 'pdf':
            if _PdfReader is None:
                raise RuntimeError("未安装 pypdf 或 PyPDF2，无法读取 PDF。")
            with open(file_path, 'rb') as f:
                pdf_reader = _PdfReader(f)
                for page in pdf_reader.pages:
                    text += page.extract_text() or ""
        elif file_type in ['doc', 'docx']:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + '\n'
    except Exception as e:
        print(f"提取文本时出错：{str(e)}")
    
    return text

def extract_knowledge_from_text(text):
    """从文本中自动提取知识，进行结构化处理"""
    if not text:
        return text
    
    import re
    
    # 增强的文档预处理：自动识别和处理不同类型的内容结构
    processed_text = text
    
    # 1. 识别并处理表格结构
    # 表格通常有固定的分隔符，如	或|或多个空格
    lines = processed_text.split('\n')
    table_lines = []
    in_table = False
    table_start_pos = 0
    
    for i, line in enumerate(lines):
        # 检测表格行：包含多个分隔符或固定宽度格式
        has_tabs = '\t' in line
        has_pipes = line.count('|') >= 2
        has_multiple_spaces = re.search(r'\s{3,}', line)
        
        if has_tabs or has_pipes or has_multiple_spaces:
            if not in_table:
                in_table = True
                table_start_pos = i
                table_lines = [line]
            else:
                table_lines.append(line)
        else:
            if in_table and table_lines:
                # 表格结束，处理整个表格
                table_content = '\n'.join(table_lines)
                if has_pipes:
                    # 处理|分隔的表格
                    table_rows = []
                    for table_line in table_lines:
                        cells = [cell.strip() for cell in table_line.split('|') if cell.strip()]
                        if cells:
                            table_rows.append('，'.join(cells))
                    if table_rows:
                        processed_table = '\n\n【表格内容】\n' + '\n'.join(table_rows) + '\n\n'
                    else:
                        processed_table = '\n\n' + table_content + '\n\n'
                else:
                    # 处理其他表格格式
                    processed_table = '\n\n【表格内容】\n' + re.sub(r'\s{2,}', ' ', table_content) + '\n\n'
                
                # 替换原始表格内容
                original_table = '\n'.join(table_lines)
                processed_text = processed_text.replace(original_table, processed_table)
                
                in_table = False
                table_lines = []
    
    # 处理最后一个表格（如果有）
    if in_table and table_lines:
        table_content = '\n'.join(table_lines)
        processed_text = processed_text.replace(table_content, '\n\n【表格内容】\n' + re.sub(r'\s{2,}', ' ', table_content) + '\n\n')
    
    # 2. 识别并处理代码块
    # 代码块通常有特定的标记或缩进
    code_block_pattern = re.compile(r'(```[\s\S]*?```|~~~[\s\S]*?~~~|(?:\n|^)[ \t]{4,}[\s\S]*?(?:\n\S|$))')
    processed_text = code_block_pattern.sub(lambda m: '\n\n【代码块】\n' + m.group(0).strip() + '\n\n', processed_text)
    
    # 3. 识别并增强命令示例
    # 识别命令行示例（如$ ls -l, > cd /home等）并添加标记
    # 增强版：支持更多命令行格式，包括行首、缩进和多行命令
    command_example_pattern = re.compile(r'((?:^|\n)[ \t]*(?:\$|>|\#|\%|\>|\<)[ \t]*[a-z]+[\s\S]*?(?=(?:^|\n)[^ \t]|$))', re.MULTILINE)
    processed_text = command_example_pattern.sub(lambda m: '\n\n【命令示例】\n' + m.group(0).strip() + '\n\n', processed_text)
    
    # 补充识别没有提示符的命令示例（如ls -l）
    bare_command_pattern = re.compile(r'(\n\s*[a-z]+\s+[-a-zA-Z0-9./]+\s*[^。；！？]*[。；！？]?)', re.MULTILINE)
    processed_text = bare_command_pattern.sub(lambda m: '\n\n【命令示例】\n' + m.group(0).strip() + '\n\n', processed_text)
    
    # 额外处理多行命令示例，确保完整匹配
    multi_line_pattern = re.compile(r'【命令示例】\n([a-z]+[\s\S]*?)(?=\n\n|$)', re.MULTILINE)
    processed_text = multi_line_pattern.sub(lambda m: '\n\n【命令示例】\n' + m.group(1).strip() + '\n\n', processed_text)
    
    # 4. 识别并提取关键概念和定义
    # 识别形如"XX是YY"的定义
    definition_patterns = [
        # 英文命令定义（如"ls 命令是..."）
        re.compile(r'([a-zA-Z]+)(?:\s*命令)?\s*(?:是|指|定义为)\s*([^。；！？]*[。；！？])'),
        # 中文命令定义（如"列表命令是..."）
        re.compile(r'([\u4e00-\u9fa5]+命令)\s*(?:是|指|定义为)\s*([^。；！？]*[。；！？])')
    ]
    
    for pattern in definition_patterns:
        processed_text = pattern.sub(lambda m: '\n\n【定义】' + m.group(1) + '：' + m.group(2) + '\n\n', processed_text)
    
    # 5. 识别并提取命令参数说明
    # 识别形如"-l 参数表示..."的参数说明
    parameter_pattern = re.compile(r'([-]+[a-zA-Z0-9]+)\s*(?:参数|选项|表示|用于)\s*([^。；！？]*[。；！？])')
    processed_text = parameter_pattern.sub(lambda m: '\n\n【参数说明】' + m.group(1) + '：' + m.group(2) + '\n\n', processed_text)
    
    # 6. 识别并提取功能描述
    # 识别形如"ls命令可以..."的功能描述
    function_pattern = re.compile(r'([a-zA-Z]+[\u4e00-\u9fa5]*[a-zA-Z]*?命令)\s*(?:可以|能够|提供)\s*([^。；！？]*[。；！？])')
    processed_text = function_pattern.sub(lambda m: '\n\n【功能说明】' + m.group(1) + '：' + m.group(2) + '\n\n', processed_text)
    
    # 7. 识别并处理章节标题
    # 章节标题通常有特定的格式（如"1. 简介"、"第一章、概述"等）
    chapter_patterns = [
        # 数字格式（如"1. 简介"、"2.1 功能"）
        re.compile(r'(^\s*\d+(?:\.\d+)*\s*[\w\W]+?\n)', re.MULTILINE),
        # 中文格式（如"第一章、概述"、"第一节 简介"）
        re.compile(r'(^\s*第[一二三四五六七八九十百千]+[章节节]、?\s*[\w\W]+?\n)', re.MULTILINE)
    ]
    
    for pattern in chapter_patterns:
        processed_text = pattern.sub(lambda m: '\n\n【章节标题】\n' + m.group(0).strip() + '\n\n', processed_text)
    
    # 8. 清理多余的空行
    processed_text = re.sub(r'\n{3,}', '\n\n', processed_text)
    
    # 9. 移除冗余的标记（避免重复标记同一内容）
    # 清理同一行或连续文本中的重复标记
    processed_text = re.sub(r'【定义】([^：]+：)+', '【定义】', processed_text)
    # 清理连续的相同标记
    processed_text = re.sub(r'(【定义】\s*)+', '【定义】', processed_text)
    # 清理重复的命令示例标记
    processed_text = re.sub(r'(【命令示例】\s*)+', '【命令示例】', processed_text)
    processed_text = re.sub(r'(【代码块】\s*)+', '【代码块】', processed_text)
    # 只清理特定的冗余标记，保留命令示例和代码块
    processed_text = re.sub(r'(【章节标题】\s*)+', '', processed_text)
    processed_text = re.sub(r'(【表格内容】\s*)+', '', processed_text)
    
    # 确保句子完整性，添加缺失的标点符号
    processed_text = re.sub(r'([^。；！？])$', r'\1。', processed_text)
    # 清理多余的空格和制表符，但保留换行符
    processed_text = re.sub(r'[ \t]+', ' ', processed_text)
    # 确保每个标记之间有换行符，这样可以被分割为不同的段落
    processed_text = re.sub(r'(【定义】|【命令示例】|【代码块】|【功能说明】|【参数说明】)', r'\n\1', processed_text)
    # 恢复正确的换行格式
    processed_text = re.sub(r'\n\s*\n', '\n\n', processed_text)
    # 清理重复的空行
    processed_text = re.sub(r'\n{3,}', '\n\n', processed_text)
    
    return processed_text

from .ai_similarity import calculate_similarity, generate_ai_content

def generate_ai_response(message, knowledge_base):
    """直接使用千问AI的接口进行作答，删除其他逻辑"""
    # 直接调用千问AI生成回答
    ai_generated = generate_ai_content(message)
    if ai_generated:
        return f"根据千问AI的回答：\n\n{ai_generated}"
    else:
        # 如果千问AI生成接口失败，返回通用回复
        return '抱歉，当前千问AI接口不可用，请稍后再试。'

def index(request):
    """首页"""
    # 获取总题数
    total_questions = Question.objects.count()
    
    if total_questions == 0:
        return render(request, 'questions/no_questions.html')
    
    # 获取第一题ID
    first_question = Question.objects.first()
    first_question_id = first_question.id if first_question else None
    
    context = {
        'total_questions': total_questions,
        'first_question_id': first_question_id
    }
    
    return render(request, 'questions/index.html', context)

def user_register(request):
    """用户注册"""
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            # 自动登录新注册的用户
            login(request, user)
            messages.success(request, '注册成功！欢迎使用在线题库系统。')
            return redirect('index')
        else:
            messages.error(request, '注册失败，请检查输入信息。')
    else:
        form = UserCreationForm()
    
    context = {
        'form': form
    }
    
    return render(request, 'questions/register.html', context)

def user_login(request):
    """用户登录"""
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            username = form.cleaned_data.get('username')
            password = form.cleaned_data.get('password')
            user = authenticate(username=username, password=password)
            if user is not None:
                login(request, user)
                messages.success(request, f'登录成功！欢迎回来，{username}。')
                # 设置session标志，用于前端清除AI聊天记录
                request.session['reset_ai_chat'] = True
                return redirect('index')
        messages.error(request, '登录失败，请检查用户名或密码。')
    else:
        form = AuthenticationForm()
    
    context = {
        'form': form
    }
    
    return render(request, 'questions/login.html', context)

def user_logout(request):
    """用户注销"""
    # 设置session标志，用于前端清除AI聊天记录
    request.session['reset_ai_chat'] = True
    logout(request)
    messages.success(request, '已成功注销登录。')
    return redirect('index')

def reset_ai_chat(request):
    """清除AI聊天记录的session标志"""
    if request.method == 'POST':
        if 'reset_ai_chat' in request.session:
            del request.session['reset_ai_chat']
    return JsonResponse({'status': 'ok'})

@login_required
def dashboard(request):
    """用户个人中心"""
    # 获取当前用户
    user = request.user
    
    # 获取所有选择题
    all_mc_questions = Question.objects.filter(type='选择题')
    total_mc_questions = all_mc_questions.count()
    
    # 获取用户已做的选择题
    answered_mc_questions = UserAnswer.objects.filter(
        user=user, 
        question__in=all_mc_questions
    )
    answered_mc_count = answered_mc_questions.count()
    
    # 计算未做的选择题数量
    unanswered_mc_count = total_mc_questions - answered_mc_count
    
    # 计算准确率
    correct_count = 0
    for user_answer in answered_mc_questions:
        if user_answer.user_answer.upper() == user_answer.question.answer.upper():
            correct_count += 1
    
    accuracy = 0
    if answered_mc_count > 0:
        accuracy = round((correct_count / answered_mc_count) * 100, 2)
    
    context = {
        'user': user,
        'total_mc_questions': total_mc_questions,
        'answered_mc_count': answered_mc_count,
        'unanswered_mc_count': unanswered_mc_count,
        'accuracy': accuracy,
        'correct_count': correct_count
    }
    
    return render(request, 'questions/dashboard.html', context)

@login_required
def start_quiz(request):
    """开始随机出题（最多10题）"""
    # 优化：只获取需要的字段，减少内存占用
    all_questions = list(Question.objects.only('id', 'type', 'question', 'options', 'answer', 'analysis'))
    
    if not all_questions:
        return redirect('index')
    
    # 随机选择最多10题
    random_questions = random.sample(all_questions, min(10, len(all_questions)))
    
    # 提取题目的ID列表
    question_ids = [q.id for q in random_questions]
    
    # 清空用户之前对这些题目的答案
    UserAnswer.objects.filter(user=request.user, question__id__in=question_ids).delete()
    
    # 初始化session变量
    request.session['quiz_questions'] = question_ids
    request.session['current_quiz_index'] = 0
    request.session['wrong_mc_indices'] = []  # 初始化答错的选择题索引列表
    request.session['correct_mc_indices'] = []  # 初始化答对的选择题索引列表
    
    # 重定向到第一题
    return redirect('answer_question', question_id=question_ids[0])

@login_required
def answer_question(request, question_id):
    """答题页面"""
    question = get_object_or_404(Question, id=question_id)
    
    # 检查是否在随机出题模式
    quiz_questions = request.session.get('quiz_questions', None)
    
    # 获取用户之前保存的答案
    user_answer = None
    similarity_score = None
    similarity_percentage = None
    if request.user.is_authenticated:
        existing_answer = UserAnswer.objects.filter(user=request.user, question=question).first()
        if existing_answer:
            user_answer = existing_answer.user_answer
            
            # 如果是简答题且有正确答案，计算相似度
            if question.type == '简答题' and question.answer:
                # 使用原有TF-IDF方法计算相似度（不使用阿里云API）
                similarity_score = fallback_calculate_similarity(user_answer, question.answer)
                
                # 计算相似度百分比
                similarity_percentage = similarity_score * 100
    if quiz_questions:
        # 随机出题模式
        current_index = quiz_questions.index(question_id) if question_id in quiz_questions else -1
        
        # 更新当前索引
        request.session['current_quiz_index'] = current_index
        
        # 获取上一题和下一题
        previous_question = None
        next_question = None
        
        if current_index > 0:
            previous_question = Question.objects.get(id=quiz_questions[current_index - 1])
        
        if current_index < len(quiz_questions) - 1:
            next_question = Question.objects.get(id=quiz_questions[current_index + 1])
        
        # 计算进度信息
        total_questions = len(quiz_questions)
        current_number = current_index + 1
    else:
        # 普通模式
        # 获取上一题和下一题
        previous_question = Question.objects.filter(id__lt=question_id).order_by('-id').first()
        next_question = Question.objects.filter(id__gt=question_id).order_by('id').first()
        
        # 获取总题数和当前题数
        total_questions = Question.objects.count()
        current_number = Question.objects.filter(id__lte=question_id).count()
        
        # 优化：只获取所有题目的ID和类型，减少内存占用
        all_questions = Question.objects.only('id', 'type')
        question_status = {}
        for q in all_questions:
            user_ans = None
            if request.user.is_authenticated:
                user_ans = UserAnswer.objects.filter(user=request.user, question=q).first()
            if user_ans and q.type == '选择题':
                if user_ans.user_answer.upper() == q.answer.upper():
                    question_status[q.id] = 'correct'
                else:
                    question_status[q.id] = 'wrong'
            elif user_ans:
                # 简答题仅标记为已回答
                question_status[q.id] = 'answered'
            else:
                question_status[q.id] = 'unanswered'
    
    # 计算进度百分比
    progress_percentage = 0
    if total_questions > 0:
        progress_percentage = round((current_number / total_questions) * 100, 2)
    
    context = {
        'question': question,
        'user_answer': user_answer,
        'previous_question': previous_question,
        'next_question': next_question,
        'total_questions': total_questions,
        'current_number': current_number,
        'progress_percentage': progress_percentage,
        'is_quiz_mode': quiz_questions is not None,  # 指示是否在随机出题模式
        'quiz_questions': quiz_questions,  # 添加quiz_questions变量
        'question_status': question_status if not quiz_questions else {},  # 仅在普通模式下传递
        'similarity_score': similarity_score,  # 添加相似度分数
        'similarity_percentage': similarity_percentage if similarity_score is not None else None,  # 添加相似度百分比
    }
    
    return render(request, 'questions/answer_question.html', context)

@login_required
def save_answer(request, question_id):
    """保存用户答案"""
    if request.method == 'POST':
        question = get_object_or_404(Question, id=question_id)
        user_answer = request.POST.get('user_answer', '').strip()
        
        # 保存或更新用户答案
        existing_answer, created = UserAnswer.objects.get_or_create(
            user=request.user,
            question=question,
            defaults={'user_answer': user_answer}
        )
        
        if not created:
            existing_answer.user_answer = user_answer
            existing_answer.save()
        
        # 检查是否答对了（仅对选择题，简答题不进行判断）
        if question.type == '选择题':
            # 仅对选择题进行答案判断
            if user_answer and user_answer.upper() == question.answer.upper():
                messages.success(request, '恭喜你！答对了！')
                # 从答错列表中移除该题目（如果存在）
                wrong_mc_indices = request.session.get('wrong_mc_indices', [])
                quiz_questions = request.session.get('quiz_questions', None)
                if quiz_questions and question_id in quiz_questions:
                    index = quiz_questions.index(question_id)
                    if index in wrong_mc_indices:
                        wrong_mc_indices.remove(index)
                        request.session['wrong_mc_indices'] = wrong_mc_indices
                    # 将该题目添加到答对列表
                    correct_mc_indices = request.session.get('correct_mc_indices', [])
                    if index not in correct_mc_indices:
                        correct_mc_indices.append(index)
                        request.session['correct_mc_indices'] = correct_mc_indices
            else:
                messages.warning(request, '回答错误！')
                # 如果是选择题答错了，记录索引并显示答案
                wrong_mc_indices = request.session.get('wrong_mc_indices', [])
                quiz_questions = request.session.get('quiz_questions', None)
                if quiz_questions and question_id in quiz_questions:
                    index = quiz_questions.index(question_id)
                    if index not in wrong_mc_indices:
                        wrong_mc_indices.append(index)
                        request.session['wrong_mc_indices'] = wrong_mc_indices
                    # 从答对列表中移除该题目（如果存在）
                    correct_mc_indices = request.session.get('correct_mc_indices', [])
                    if index in correct_mc_indices:
                        correct_mc_indices.remove(index)
                        request.session['correct_mc_indices'] = correct_mc_indices
                # 显示答案
                return redirect('show_answer', question_id=question_id)
        else:
            # 简答题使用AI相似度对比判断
            if user_answer and question.answer:
                # 使用原有TF-IDF方法计算相似度（不使用阿里云API）
                similarity = fallback_calculate_similarity(user_answer, question.answer)
                
                # 计算相似度百分比
                similarity_percentage = similarity * 100
                
                # 设置相似度阈值
                similarity_threshold = 0.7
                
                # 根据相似度判断答案是否正确
                if similarity >= similarity_threshold:
                    messages.success(request, f'答案已保存！相似度分析：{round(similarity*100, 1)}%，恭喜你回答正确！')
                else:
                    messages.warning(request, f'答案已保存！相似度分析：{round(similarity*100, 1)}%，你的回答与正确答案有一定差距。可以查看答案了解详细内容。')
            else:
                messages.success(request, '答案已保存')
        
        # 检查是否在随机出题模式
        quiz_questions = request.session.get('quiz_questions', None)
        if quiz_questions:
            # 随机出题模式
            current_index = request.session.get('current_quiz_index', 0)
            
            # 获取下一题
            if current_index < len(quiz_questions) - 1:
                next_question_id = quiz_questions[current_index + 1]
                # 更新当前索引
                request.session['current_quiz_index'] = current_index + 1
                return redirect('answer_question', question_id=next_question_id)
            else:
                # 如果是最后一题，返回当前页面
                return redirect('answer_question', question_id=question_id)
        else:
            # 普通模式
            # 获取下一题，如果有的话
            next_question = Question.objects.filter(id__gt=question_id).order_by('id').first()
            if next_question:
                return redirect('answer_question', question_id=next_question.id)
            else:
                # 如果是最后一题，返回当前页面
                return redirect('answer_question', question_id=question_id)
    
    return redirect('answer_question', question_id=question_id)

@login_required
def show_answer(request, question_id):
    """显示正确答案"""
    question = get_object_or_404(Question, id=question_id)
    
    # 获取用户答案
    user_answer = None
    if request.user.is_authenticated:
        existing_answer = UserAnswer.objects.filter(user=request.user, question=question).first()
        if existing_answer:
            user_answer = existing_answer.user_answer
    
    # 获取总题数
    total_questions = Question.objects.count()
    
    # 检查是否在随机出题模式
    quiz_questions = request.session.get('quiz_questions', None)
    if quiz_questions:
        # 随机出题模式
        current_index = quiz_questions.index(question_id) if question_id in quiz_questions else -1
        
        # 获取下一题
        next_question = None
        if current_index < len(quiz_questions) - 1:
            next_question = Question.objects.get(id=quiz_questions[current_index + 1])
    else:
        # 普通模式
        # 获取下一题
        next_question = Question.objects.filter(id__gt=question_id).order_by('id').first()
        
        # 优化：只获取所有题目的ID和类型，减少内存占用
        all_questions = Question.objects.only('id', 'type')
        question_status = {}
        for q in all_questions:
            user_ans = None
            if request.user.is_authenticated:
                user_ans = UserAnswer.objects.filter(user=request.user, question=q).first()
            if user_ans and q.type == '选择题':
                if user_ans.user_answer.upper() == q.answer.upper():
                    question_status[q.id] = 'correct'
                else:
                    question_status[q.id] = 'wrong'
            elif user_ans:
                question_status[q.id] = 'answered'
            else:
                question_status[q.id] = 'unanswered'
    
    context = {
        'question': question,
        'user_answer': user_answer,
        'total_questions': total_questions,
        'quiz_questions': quiz_questions,
        'next_question': next_question,
        'question_status': question_status if not quiz_questions else {},  # 仅在普通模式下传递
    }
    
    return render(request, 'questions/show_answer.html', context)



@login_required
def add_question(request):
    """添加题目"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    if request.method == 'POST':
        # 获取表单数据
        question_type = request.POST.get('type')
        question_content = request.POST.get('question')
        answer = request.POST.get('answer')
        analysis = request.POST.get('analysis')
        
        # 获取选项（如果是选择题）
        options = []
        if question_type == '选择题':
            for i in range(1, 5):  # 最多4个选项
                option = request.POST.get(f'option{i}')
                if option:
                    options.append(option)
        
        # 获取新题目的ID（最大ID+1）
        max_id = Question.objects.aggregate(max_id=models.Max('id'))['max_id']
        new_id = max_id + 1 if max_id else 1
        
        # 创建新题目
        try:
            new_question = Question(
                id=new_id,
                type=question_type,
                question=question_content,
                answer=answer,
                options=options,
                analysis=analysis
            )
            new_question.save()
            messages.success(request, '题目添加成功！')
            return redirect('admin_dashboard')
        except Exception as e:
            messages.error(request, f'添加题目失败：{str(e)}')
            return redirect('add_question')
    
    return render(request, 'questions/add_question.html')

@login_required
def edit_question(request, question_id):
    """编辑题目"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取要编辑的题目
    question = get_object_or_404(Question, id=question_id)
    
    if request.method == 'POST':
        # 获取表单数据
        question_type = request.POST.get('type')
        question_content = request.POST.get('question')
        answer = request.POST.get('answer')
        analysis = request.POST.get('analysis')
        
        # 获取选项（如果是选择题）
        options = []
        if question_type == '选择题':
            for i in range(1, 5):  # 最多4个选项
                option = request.POST.get(f'option{i}')
                if option:
                    options.append(option)
        
        # 更新题目
        try:
            question.type = question_type
            question.question = question_content
            question.answer = answer
            question.options = options
            question.analysis = analysis
            question.save()
            messages.success(request, '题目修改成功！')
            return redirect('admin_dashboard')
        except Exception as e:
            messages.error(request, f'修改题目失败：{str(e)}')
            return redirect('edit_question', question_id=question_id)
    
    context = {
        'question': question
    }
    
    return render(request, 'questions/edit_question.html', context)

@login_required
def delete_question(request, question_id):
    """删除题目"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取要删除的题目
    question = get_object_or_404(Question, id=question_id)
    
    if request.method == 'POST':
        # 删除题目
        try:
            question.delete()
            messages.success(request, '题目删除成功！')
            return redirect('admin_dashboard')
        except Exception as e:
            messages.error(request, f'删除题目失败：{str(e)}')
            return redirect('admin_dashboard')
    
    return redirect('admin_dashboard')

@login_required
def admin_dashboard(request):
    """管理员仪表盘"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取所有用户（排除管理员）
    users = User.objects.filter(is_superuser=False)
    
    # 获取所有题目
    questions = Question.objects.all()
    
    # 计算选择题数量
    mc_questions_count = questions.filter(type='选择题').count()
    
    # 为每个用户获取AI使用信息
    user_ai_info = {}
    for user in users:
        ai_usage, created = UserAIUsage.objects.get_or_create(user=user)
        user_ai_info[user.id] = {
            'is_ai_unlocked': ai_usage.is_ai_unlocked,
            'ai_usage_count': ai_usage.ai_usage_count
        }
    
    context = {
        'users': users,
        'questions': questions,
        'mc_questions_count': mc_questions_count,
        'user_ai_info': user_ai_info
    }
    
    return render(request, 'questions/admin_dashboard.html', context)

@login_required
def view_all_questions(request):
    """查看所有题目列表"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取所有题目，按ID排序
    questions = Question.objects.all().order_by('id')
    
    context = {
        'questions': questions
    }
    
    return render(request, 'questions/all_questions.html', context)

@login_required
def view_user_stats(request, user_id):
    """查看用户答题统计"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取要查看的用户
    user = get_object_or_404(User, id=user_id)
    
    # 获取所有选择题
    all_mc_questions = Question.objects.filter(type='选择题')
    total_mc_questions = all_mc_questions.count()
    
    # 获取用户已做的选择题
    answered_mc_questions = UserAnswer.objects.filter(
        user=user, 
        question__in=all_mc_questions
    )
    answered_mc_count = answered_mc_questions.count()
    
    # 计算未做的选择题数量
    unanswered_mc_count = total_mc_questions - answered_mc_count
    
    # 计算准确率
    correct_count = 0
    for user_answer in answered_mc_questions:
        if user_answer.user_answer.upper() == user_answer.question.answer.upper():
            correct_count += 1
    
    accuracy = 0
    if answered_mc_count > 0:
        accuracy = round((correct_count / answered_mc_count) * 100, 2)
    
    # 计算题目完成度百分比
    completion_rate = 0
    completion_rate_int = 0
    if total_mc_questions > 0:
        completion_rate = round((answered_mc_count / total_mc_questions) * 100, 2)
        completion_rate_int = round(completion_rate)
    
    # 计算错误答案数量
    incorrect_count = answered_mc_count - correct_count
    
    context = {
        'viewed_user': user,
        'total_mc_questions': total_mc_questions,
        'answered_mc_count': answered_mc_count,
        'unanswered_mc_count': unanswered_mc_count,
        'accuracy': accuracy,
        'correct_count': correct_count,
        'incorrect_count': incorrect_count,
        'completion_rate': completion_rate,
        'completion_rate_int': completion_rate_int
    }
    
    return render(request, 'questions/user_stats.html', context)

@login_required
def delete_user(request, user_id):
    """删除用户"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取要删除的用户
    user = get_object_or_404(User, id=user_id)
    
    # 不能删除自己
    if user.id == request.user.id:
        messages.error(request, '您不能删除自己的账号。')
        return redirect('admin_dashboard')
    
    if request.method == 'POST':
        # 删除用户
        try:
            user.delete()
            messages.success(request, '用户删除成功！')
            return redirect('admin_dashboard')
        except Exception as e:
            messages.error(request, f'删除用户失败：{str(e)}')
            return redirect('admin_dashboard')
    
    return redirect('admin_dashboard')

@login_required
def ai_document_management(request):
    """AI文档管理"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    if request.method == 'POST':
        # 获取表单数据
        title = request.POST.get('title')
        file = request.FILES.get('file')
        
        if not title or not file:
            messages.error(request, '文档标题和文件不能为空。')
            return redirect('ai_document_management')
        
        # 获取文件类型
        file_type = file.name.split('.')[-1].lower()
        if file_type not in ['pdf', 'txt', 'doc', 'docx']:
            messages.error(request, '只支持PDF、TXT、Word格式文件。')
            return redirect('ai_document_management')
        
        try:
            # 创建文档记录
            ai_document = AIDocument(
                title=title,
                file=file,
                uploaded_by=request.user,
                file_type=file_type
            )
            ai_document.save()
            
            # 提取文本内容
            file_path = ai_document.file.path
            extracted_text = extract_text_from_file(file_path, file_type)
            
            # 更新文档的文本内容
            ai_document.content = extracted_text
            ai_document.save()
            
            # 清除AI知识库缓存
            cache.delete('ai_knowledge_base')
            
            messages.success(request, '文档上传成功！文本内容已提取。')
            return redirect('ai_document_management')
        except Exception as e:
            messages.error(request, f'文档上传失败：{str(e)}')
            return redirect('ai_document_management')
    
    # 获取所有文档
    documents = AIDocument.objects.all().order_by('-uploaded_at')
    
    context = {
        'documents': documents
    }
    
    return render(request, 'questions/ai_document_management.html', context)


@login_required
def bookshelf(request):
    """公共书架 - 查看所有管理员上传的文档"""
    # 获取所有管理员上传的文档
    documents = AIDocument.objects.all().order_by('-uploaded_at')
    context = {'documents': documents}
    return render(request, 'questions/bookshelf.html', context)


@login_required
def delete_ai_document(request, document_id):
    """删除AI文档"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('index')
    
    # 获取要删除的文档
    document = get_object_or_404(AIDocument, id=document_id)
    
    if request.method == 'POST':
        # 删除文档文件
        if os.path.exists(document.file.path):
            os.remove(document.file.path)
        
        # 删除文档记录
        document.delete()
        
        # 清除AI知识库缓存
        cache.delete('ai_knowledge_base')
        
        messages.success(request, '文档删除成功！')
        return redirect('ai_document_management')
    
    return redirect('ai_document_management')

from django.views.decorators.csrf import csrf_exempt

@csrf_exempt
def ai_chat(request):
    """AI聊天API"""
    if request.method == 'POST':
        try:
            # 尝试解析JSON数据
            data = json.loads(request.body)
            message = data.get('message', '').strip()
        except json.JSONDecodeError:
            # 如果JSON解析失败，尝试从表单中获取
            message = request.POST.get('message', '').strip()
        
        if not message:
            return JsonResponse({'error': '消息不能为空'}, status=400)
        
        # 检查用户AI使用限制
        if request.user.is_authenticated:
            # 获取或创建用户AI使用记录
            user_ai_usage, created = UserAIUsage.objects.get_or_create(user=request.user)
            
            # 检查是否有使用权限
            if not request.user.is_superuser and not user_ai_usage.is_ai_unlocked:
                if user_ai_usage.ai_usage_count >= user_ai_usage.max_ai_usage:
                    return JsonResponse({'response': '您的AI使用次数已达上限，请联系管理员解禁AI功能。'})
        
        # 优化：使用缓存存储知识库内容，避免每次查询数据库
        cache_key = 'ai_knowledge_base'
        knowledge_base = cache.get(cache_key)
        
        if knowledge_base is None:
            # 缓存未命中，从数据库加载
            documents = AIDocument.objects.only('title', 'content')
            knowledge_base = ""
            
            for doc in documents:
                if doc.content:
                    knowledge_base += f"\n\n--- {doc.title} ---\n{doc.content}"
            
            # 缓存知识库内容，缓存时间设置为1小时（3600秒）
            cache.set(cache_key, knowledge_base, 3600)
        
        # 生成AI回复
        ai_response = generate_ai_response(message, knowledge_base)
        
        # 保存对话记录（仅当用户登录时）
        if request.user.is_authenticated:
            AIChat.objects.create(
                user=request.user,
                message=message,
                response=ai_response
            )
            
            # 增加用户的AI使用次数（如果用户不是超级管理员且未被解禁）
            if not request.user.is_superuser:
                user_ai_usage, created = UserAIUsage.objects.get_or_create(user=request.user)
                if not user_ai_usage.is_ai_unlocked:
                    user_ai_usage.ai_usage_count += 1
                    user_ai_usage.save()
        
        return JsonResponse({'response': ai_response})
    
    return JsonResponse({'error': '无效请求'}, status=400)

@login_required
def unlock_ai(request, user_id):
    """管理员解禁/锁定用户AI功能"""
    # 检查是否为管理员
    if not request.user.is_superuser:
        messages.error(request, '您没有权限执行此操作。')
        return redirect('admin_dashboard')
    
    # 获取目标用户
    target_user = get_object_or_404(User, id=user_id)
    
    # 获取或创建用户AI使用记录
    user_ai_usage, created = UserAIUsage.objects.get_or_create(user=target_user)
    
    # 切换解禁状态
    user_ai_usage.is_ai_unlocked = not user_ai_usage.is_ai_unlocked
    user_ai_usage.save()
    
    # 添加成功消息
    if user_ai_usage.is_ai_unlocked:
        messages.success(request, f'已成功解禁用户{target_user.username}的AI功能！')
    else:
        messages.success(request, f'已成功锁定用户{target_user.username}的AI功能！')
    
    # 返回用户列表页面
    return redirect('admin_dashboard')


def get_chat_history(request):
    """获取当前用户最近的聊天记录"""
    if not request.user.is_authenticated:
        return JsonResponse({'error': '未登录'}, status=401)
    
    # 获取当前用户最近的15条聊天记录
    chats = AIChat.objects.filter(user=request.user).order_by('-timestamp')[:15]
    
    # 转换为前端需要的格式
    messages = []
    for chat in reversed(chats):  # 反转顺序，按时间正序排列
        messages.append({
            'sender': 'user',
            'content': chat.message
        })
        messages.append({
            'sender': 'ai',
            'content': chat.response
        })
    
    return JsonResponse({'messages': messages})


# ============================================================
# 模拟面试（v1：不依赖大模型，纯题库驱动 + 自动判分）
# 说明：AI 点评层通过 settings.INTERVIEW_AI_ENABLED 控制，
#       额度恢复后把它打开并实现 generate_interview_review 即可无缝接入。
# ============================================================

# 简答题自动判分阈值（基于 TF-IDF 相似度）
INTERVIEW_SA_THRESHOLD = 0.5


def _interview_pool(mode):
    """按模式取候选题目"""
    qs = Question.objects.all()
    if mode == 'mc':
        qs = qs.filter(type='选择题')
    elif mode == 'sa':
        qs = qs.filter(type='简答题')
    return qs


def _interview_ai_enabled():
    return bool(getattr(settings, 'INTERVIEW_AI_ENABLED', False))


@login_required
def interview_home(request):
    """模拟面试：开始页"""
    if request.method == 'POST':
        mode = request.POST.get('mode', 'all')
        if mode not in ('all', 'mc', 'sa'):
            mode = 'all'
        try:
            num = int(request.POST.get('num', 10))
        except (TypeError, ValueError):
            num = 10
        num = max(5, min(num, 30))

        pool = list(_interview_pool(mode).only('id'))
        if not pool:
            messages.error(request, '该模式暂无题目，请更换模式后重试。')
            return redirect('interview_home')

        picked = random.sample(pool, min(num, len(pool)))
        sess = InterviewSession.objects.create(
            user=request.user,
            mode=mode,
            question_ids=[q.id for q in picked],
            total=len(picked),
        )
        return redirect('interview_run', session_id=sess.id)

    counts = {
        'all': Question.objects.count(),
        'mc': Question.objects.filter(type='选择题').count(),
        'sa': Question.objects.filter(type='简答题').count(),
    }
    recent = InterviewSession.objects.filter(user=request.user, is_finished=True)[:5]
    return render(request, 'questions/interview_home.html', {
        'counts': counts,
        'recent': recent,
    })


def _interview_valid_ids(sess):
    """返回题库中仍然存在的题目 id（保持出题顺序）"""
    ids = list(sess.question_ids or [])
    existing = set(Question.objects.filter(id__in=ids).values_list('id', flat=True))
    return [i for i in ids if i in existing]


def _interview_question_payload(q):
    return {
        'id': q.id,
        'type': q.type,
        'question': q.question,
        'options': q.options or [],
    }


def _interview_time_limit(qtype):
    """每题限时（秒），可在 settings 覆盖：选择题默认 45s、简答题默认 120s"""
    if qtype == '简答题':
        return int(getattr(settings, 'INTERVIEW_TIME_SA', 120))
    return int(getattr(settings, 'INTERVIEW_TIME_MC', 45))


def _interview_answered_count(sess):
    return InterviewAnswer.objects.filter(session=sess).count()


def _interview_grade(sess):
    """按已保存的逐题作答判分（未作答记错），并结算整场面试"""
    ids = _interview_valid_ids(sess)
    qmap = {q.id: q for q in Question.objects.filter(id__in=ids)}
    recorded = {a.order: a for a in sess.answers.all()}

    correct = 0
    graded_total = 0
    for idx, qid in enumerate(ids):
        q = qmap.get(qid)
        if not q:
            continue
        graded_total += 1
        rec = recorded.get(idx)
        ua = (rec.user_answer or '').strip() if rec else ''
        is_correct = False
        sim = 0.0
        if q.type == '选择题':
            is_correct = bool(ua) and ua.upper() == (q.answer or '').strip().upper()
        else:
            if ua:
                try:
                    sim = float(fallback_calculate_similarity(ua, q.answer or ''))
                except Exception:
                    sim = 0.0
                is_correct = sim >= INTERVIEW_SA_THRESHOLD
        if rec:
            rec.is_correct = is_correct
            rec.similarity = round(sim, 4)
            rec.save(update_fields=['is_correct', 'similarity'])
        else:
            InterviewAnswer.objects.create(
                session=sess, question=q, order=idx,
                user_answer='', is_correct=is_correct, similarity=round(sim, 4),
            )
        if is_correct:
            correct += 1

    sess.total = graded_total
    sess.correct = correct
    sess.score = round(correct / graded_total * 100) if graded_total else 0
    sess.is_finished = True
    sess.finished_at = timezone.now()
    if sess.started_at:
        sess.duration_seconds = max(0, int((sess.finished_at - sess.started_at).total_seconds()))
    sess.save()


@login_required
def interview_run(request, session_id):
    """模拟面试：逐题作答页（单题沉浸 / 每题限时 / 交卷前不可回看）"""
    sess = get_object_or_404(InterviewSession, id=session_id, user=request.user)
    if sess.is_finished:
        return redirect('interview_report', session_id=sess.id)

    ids = _interview_valid_ids(sess)
    total = len(ids)
    answered = _interview_answered_count(sess)

    # 题目已全部作答（如答完最后一题后刷新）→ 直接结算
    if total == 0 or answered >= total:
        _interview_grade(sess)
        return redirect('interview_report', session_id=sess.id)

    current_q = Question.objects.get(id=ids[answered])
    return render(request, 'questions/interview_run.html', {
        'session': sess,
        'index': answered + 1,
        'total': total,
        'question_json': json.dumps(_interview_question_payload(current_q), ensure_ascii=False),
        'limit': _interview_time_limit(current_q.type),
    })


@login_required
def interview_answer(request, session_id):
    """模拟面试：提交当前题的作答（服务端逐题推进，禁止回看/修改）"""
    sess = get_object_or_404(InterviewSession, id=session_id, user=request.user)
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'method_not_allowed'}, status=405)
    if sess.is_finished:
        return JsonResponse({'ok': True, 'done': True,
                             'redirect': reverse('interview_report', args=[sess.id])})

    try:
        data = json.loads(request.body or b'{}')
    except (json.JSONDecodeError, TypeError):
        data = {}
    answer = str(data.get('answer', '') or '').strip()

    ids = _interview_valid_ids(sess)
    idx = _interview_answered_count(sess)
    if idx >= len(ids):
        return JsonResponse({'ok': True, 'done': True})

    q = Question.objects.get(id=ids[idx])
    InterviewAnswer.objects.create(
        session=sess, question=q, order=idx,
        user_answer=answer, is_correct=False, similarity=0,
    )

    nxt = idx + 1
    if nxt >= len(ids):
        return JsonResponse({'ok': True, 'done': True})

    nx = Question.objects.get(id=ids[nxt])
    return JsonResponse({
        'ok': True,
        'done': False,
        'index': nxt + 1,
        'total': len(ids),
        'question': _interview_question_payload(nx),
        'limit': _interview_time_limit(nx.type),
    })


@login_required
def interview_submit(request, session_id):
    """模拟面试：交卷并自动判分（按服务端保存的逐题作答结算）"""
    sess = get_object_or_404(InterviewSession, id=session_id, user=request.user)
    if request.method != 'POST':
        return redirect('interview_run', session_id=sess.id)
    if sess.is_finished:
        return JsonResponse({'ok': True, 'redirect': reverse('interview_report', args=[sess.id])})

    _interview_grade(sess)
    return JsonResponse({'ok': True, 'redirect': reverse('interview_report', args=[sess.id])})


def generate_interview_review(sess):
    """AI 点评层（预留）。

    额度恢复后：把 settings.INTERVIEW_AI_ENABLED 设为 True，
    并在本函数里调用 generate_ai_content(...) 生成针对本场面试的点评/追问。
    目前返回 None，报告页会自动隐藏该区块。
    """
    if not _interview_ai_enabled():
        return None
    # 预留实现位（当前不依赖大模型，故不做任何调用）
    return None


@login_required
def interview_report(request, session_id):
    """模拟面试：成绩报告页"""
    sess = get_object_or_404(InterviewSession, id=session_id, user=request.user)
    if not sess.is_finished:
        return redirect('interview_run', session_id=sess.id)

    answers = list(sess.answers.select_related('question').order_by('order'))
    context = {
        'session': sess,
        'answers': answers,
        'ai_enabled': _interview_ai_enabled(),
        'ai_review': generate_interview_review(sess),
    }
    return render(request, 'questions/interview_report.html', context)


@login_required
def interview_history(request):
    """模拟面试：历史记录"""
    sessions = InterviewSession.objects.filter(user=request.user).order_by('-started_at')[:50]
    return render(request, 'questions/interview_history.html', {'sessions': sessions})


# ============================================================
# K8s 学习中心（搬运自静态站点 k8s-quest）
# 文件存放在 questions/k8s_quest/ 下，由本视图托管：
#   /k8s-learning/        → index.html
#   /k8s-learning/<路径>   → 其余静态资源（css/js/章节/排行榜/我的）
# 全部要求登录（沿用题库系统用户认证），并在 HTML 中注入当前登录用户。
# ============================================================
K8S_QUEST_ROOT = os.path.join(settings.BASE_DIR, 'questions', 'k8s_quest')


def _k8s_inject_user(html, request):
    """在静态 HTML 的 </head> 前注入：题库系统登录用户 + 返回题库入口。"""
    user = request.user
    nickname = user.get_full_name() or user.username
    payload = json.dumps({'username': user.username, 'nickname': nickname}, ensure_ascii=False)
    bridge = (
        '<script>'
        'window.K8S_USER=' + payload + ';'
        'window.K8S_LOGOUT_URL=' + json.dumps(reverse('logout')) + ';'
        '(function(){'
        'function add(){'
        'if(document.getElementById("k8s-back")){return;}'
        'var a=document.createElement("a");a.id="k8s-back";'
        'a.href=' + json.dumps(reverse('index')) + ';'
        'a.textContent="← 返回题库";'
        'a.style.cssText="position:fixed;left:16px;bottom:16px;z-index:9998;'
        'background:#326ce5;color:#fff;padding:8px 14px;border-radius:999px;'
        'font-size:.85rem;text-decoration:none;box-shadow:0 4px 14px rgba(0,0,0,.25);";'
        'document.body.appendChild(a);}'
        'if(document.readyState!=="loading"){add();}'
        'else{document.addEventListener("DOMContentLoaded",add);}'
        '})();'
        '</script>'
    )
    idx = html.lower().rfind('</head>')
    if idx == -1:
        return html + bridge
    return html[:idx] + bridge + html[idx:]


@login_required
def k8s_learning(request, asset_path=''):
    """托管 K8s 学习中心静态站点（需登录），HTML 页面注入当前用户。"""
    target = os.path.normpath(os.path.join(K8S_QUEST_ROOT, asset_path or 'index.html'))
    # 防目录穿越：解析后的路径必须仍在站点根目录内
    if target != K8S_QUEST_ROOT and not target.startswith(K8S_QUEST_ROOT + os.sep):
        raise Http404('Not found')
    if os.path.isdir(target):
        target = os.path.join(target, 'index.html')
    if not os.path.isfile(target):
        raise Http404('Not found')

    content_type = mimetypes.guess_type(target)[0] or 'application/octet-stream'
    if content_type.startswith('text/html'):
        with open(target, 'r', encoding='utf-8') as fh:
            html = fh.read()
        return HttpResponse(_k8s_inject_user(html, request),
                            content_type='text/html; charset=utf-8')
    return FileResponse(open(target, 'rb'), content_type=content_type)


# ============================================================
# K8s 学习中心服务端接口（进度持久化 + 全站共享排行榜）
# 供 k8s-quest 站点的 game.js / leaderboard.html 调用：
#   GET  /api/leaderboard?limit=N  -> [{user_id, nickname, stars}]
#   GET  /api/progress             -> {stars, chapters, nickname}
#   POST /api/submit               -> {chapterId, questionId, starCount, totalQuestions}
#   POST /api/spend                -> {count}
# 认证沿用题库系统的登录会话（request.user），无需额外 token。
# ============================================================
def _k8s_api_required(view_func):
    """接口统一要求登录：未登录返回 401 JSON，避免被重定向到登录页。"""
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return JsonResponse({'error': 'login required'}, status=401)
        return view_func(request, *args, **kwargs)
    return _wrapped


def _k8s_nickname(user):
    return user.get_full_name() or user.username


def _k8s_earned_stars(prog):
    """由章节答题进度派生「本中心答题所获」星星数。

    只统计通过答题获得的星星（题目分 + 章节通关奖励），
    不扣消费、不含历史脏数据，与前端本地加星口径保持一致。"""
    chapters = prog.chapters or {}
    total = 0
    for ch in chapters.values():
        if not isinstance(ch, dict):
            continue
        try:
            total += int(ch.get('stars', 0) or 0)
        except (TypeError, ValueError):
            pass
        if ch.get('completed'):
            total += 5  # 章节通关奖励
    return total


def _k8s_progress_data(user, prog=None):
    """返回接口/前端统一使用的进度结构。"""
    if prog is None:
        prog, _ = K8sProgress.objects.get_or_create(user=user)
    return {
        'stars': _k8s_earned_stars(prog),
        'chapters': prog.chapters or {},
        'nickname': _k8s_nickname(user),
    }


@csrf_exempt
@_k8s_api_required
def k8s_api_progress(request):
    """读取当前登录用户的 K8s 学习进度。"""
    return JsonResponse(_k8s_progress_data(request.user))


@csrf_exempt
@_k8s_api_required
def k8s_api_leaderboard(request):
    """全站共享排行榜：按星星数降序返回所有用户。"""
    try:
        limit = int(request.GET.get('limit', 50))
    except (TypeError, ValueError):
        limit = 50
    limit = max(1, min(limit, 200))
    rows = list(K8sProgress.objects.select_related('user'))
    scored = [(_k8s_earned_stars(row), row) for row in rows]
    scored.sort(key=lambda item: (-item[0], item[1].updated_at))
    data = [{
        'user_id': row.user.username,
        'nickname': _k8s_nickname(row.user),
        'stars': stars,
    } for stars, row in scored[:limit]]
    return JsonResponse(data, safe=False)


@csrf_exempt
@_k8s_api_required
def k8s_api_submit(request):
    """提交一道题目的作答结果，加星并持久化到服务器。"""
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)
    try:
        body = json.loads(request.body or b'{}')
    except json.JSONDecodeError:
        body = {}

    chapter_id = str(body.get('chapterId') or '').strip()
    question_id = str(body.get('questionId') or '').strip()
    if not chapter_id or not question_id:
        return JsonResponse({'error': 'invalid payload'}, status=400)

    try:
        star_count = int(body.get('starCount') or 0)
    except (TypeError, ValueError):
        star_count = 0
    try:
        total_q = int(body.get('totalQuestions') or 0)
    except (TypeError, ValueError):
        total_q = 0

    prog, _ = K8sProgress.objects.get_or_create(user=request.user)
    chapters = prog.chapters or {}
    ch = chapters.get(chapter_id) or {'completed': False, 'stars': 0, 'questions': {}}
    questions = ch.get('questions') or {}

    # 同一题只计一次分，重复提交不加星
    if questions.get(question_id):
        return JsonResponse({'alreadyDone': True, 'data': _k8s_progress_data(request.user, prog)})

    questions[question_id] = True
    ch['questions'] = questions
    ch['stars'] = int(ch.get('stars', 0)) + star_count
    prog.stars += star_count

    chapter_completed = False
    if not ch.get('completed'):
        denom = total_q or len(questions) or 1
        if len(questions) / denom >= 0.8:
            ch['completed'] = True
            prog.stars += 5  # 章节通关奖励，与前端本地逻辑保持一致
            chapter_completed = True

    chapters[chapter_id] = ch
    prog.chapters = chapters
    prog.save(update_fields=['stars', 'chapters', 'updated_at'])

    new_stars = star_count + (5 if chapter_completed else 0)
    payload = _k8s_progress_data(request.user, prog)
    return JsonResponse({
        'alreadyDone': False,
        'newStars': new_stars,
        'chapterCompleted': chapter_completed,
        'data': payload,
    })


@csrf_exempt
@_k8s_api_required
def k8s_api_spend(request):
    """消耗星星（答错换题等），星星不足时返回 success=false。"""
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)
    try:
        body = json.loads(request.body or b'{}')
    except json.JSONDecodeError:
        body = {}
    try:
        count = int(body.get('count') or 1)
    except (TypeError, ValueError):
        count = 1
    count = max(0, count)

    prog, _ = K8sProgress.objects.get_or_create(user=request.user)
    data = _k8s_progress_data(request.user, prog)
    if prog.stars < count:
        return JsonResponse({'success': False, 'reason': 'stars', 'data': data})

    prog.stars -= count
    prog.save(update_fields=['stars', 'updated_at'])
    data = _k8s_progress_data(request.user, prog)
    return JsonResponse({'success': True, 'data': data})


# ============================================================
# LangChain 学习中心（搬运自静态站点 langchain-quest）
# 文件存放在 questions/langchain_quest/ 下，由本视图托管：
#   /langchain-learning/        → index.html
#   /langchain-learning/<路径>   → 其余静态资源（css/js/章节/排行榜/我的）
# 全部要求登录（沿用题库系统用户认证），并在 HTML 中注入当前登录用户。
# ============================================================
LANGCHAIN_QUEST_ROOT = os.path.join(settings.BASE_DIR, 'questions', 'langchain_quest')


def _langchain_inject_user(html, request):
    """在静态 HTML 的 </head> 前注入：题库系统登录用户 + 返回题库入口。"""
    user = request.user
    nickname = user.get_full_name() or user.username
    payload = json.dumps({'username': user.username, 'nickname': nickname}, ensure_ascii=False)
    bridge = (
        '<script>'
        'window.LANGCHAIN_USER=' + payload + ';'
        'window.LANGCHAIN_LOGOUT_URL=' + json.dumps(reverse('logout')) + ';'
        '(function(){'
        'function add(){'
        'if(document.getElementById("langchain-back")){return;}'
        'var a=document.createElement("a");a.id="langchain-back";'
        'a.href=' + json.dumps(reverse('index')) + ';'
        'a.textContent="← 返回题库";'
        'a.style.cssText="position:fixed;left:16px;bottom:16px;z-index:9998;'
        'background:#00a3c4;color:#fff;padding:8px 14px;border-radius:999px;'
        'font-size:.85rem;text-decoration:none;box-shadow:0 4px 14px rgba(0,0,0,.25);";'
        'document.body.appendChild(a);}'
        'if(document.readyState!=="loading"){add();}'
        'else{document.addEventListener("DOMContentLoaded",add);}'
        '})();'
        '</script>'
    )
    idx = html.lower().rfind('</head>')
    if idx == -1:
        return html + bridge
    return html[:idx] + bridge + html[idx:]


@login_required
def langchain_learning(request, asset_path=''):
    """托管 LangChain 学习中心静态站点（需登录），HTML 页面注入当前用户。"""
    target = os.path.normpath(os.path.join(LANGCHAIN_QUEST_ROOT, asset_path or 'index.html'))
    # 防目录穿越：解析后的路径必须仍在站点根目录内
    if target != LANGCHAIN_QUEST_ROOT and not target.startswith(LANGCHAIN_QUEST_ROOT + os.sep):
        raise Http404('Not found')
    if os.path.isdir(target):
        target = os.path.join(target, 'index.html')
    if not os.path.isfile(target):
        raise Http404('Not found')

    content_type = mimetypes.guess_type(target)[0] or 'application/octet-stream'
    if content_type.startswith('text/html'):
        with open(target, 'r', encoding='utf-8') as fh:
            html = fh.read()
        return HttpResponse(_langchain_inject_user(html, request),
                            content_type='text/html; charset=utf-8')
    return FileResponse(open(target, 'rb'), content_type=content_type)


# ============================================================
# LangChain 学习中心服务端接口（进度持久化 + 全站共享排行榜）
# 供 langchain-quest 站点的 game.js / leaderboard.html 调用：
#   GET  /langchain-api/leaderboard?limit=N  -> [{user_id, nickname, stars}]
#   GET  /langchain-api/progress             -> {stars, chapters, nickname}
#   POST /langchain-api/submit               -> {chapterId, questionId, starCount, totalQuestions}
#   POST /langchain-api/spend                -> {count}
# 认证沿用题库系统的登录会话（request.user），无需额外 token。
# ============================================================
def _langchain_api_required(view_func):
    """接口统一要求登录：未登录返回 401 JSON，避免被重定向到登录页。"""
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return JsonResponse({'error': 'login required'}, status=401)
        return view_func(request, *args, **kwargs)
    return _wrapped


def _langchain_nickname(user):
    return user.get_full_name() or user.username


def _langchain_progress_data(user, prog=None):
    """返回接口/前端统一使用的进度结构。"""
    if prog is None:
        prog, _ = LangChainProgress.objects.get_or_create(user=user)
    return {
        'stars': prog.stars,
        'chapters': prog.chapters or {},
        'nickname': _langchain_nickname(user),
    }


@csrf_exempt
@_langchain_api_required
def langchain_api_progress(request):
    """读取当前登录用户的 LangChain 学习进度。"""
    return JsonResponse(_langchain_progress_data(request.user))


@csrf_exempt
@_langchain_api_required
def langchain_api_leaderboard(request):
    """全站共享排行榜：按星星数降序返回所有用户。"""
    try:
        limit = int(request.GET.get('limit', 50))
    except (TypeError, ValueError):
        limit = 50
    limit = max(1, min(limit, 200))
    rows = LangChainProgress.objects.select_related('user').order_by('-stars', 'updated_at')[:limit]
    data = [{
        'user_id': row.user.username,
        'nickname': _langchain_nickname(row.user),
        'stars': row.stars,
    } for row in rows]
    return JsonResponse(data, safe=False)


@csrf_exempt
@_langchain_api_required
def langchain_api_submit(request):
    """提交一道题目的作答结果，加星并持久化到服务器。"""
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)
    try:
        body = json.loads(request.body or b'{}')
    except json.JSONDecodeError:
        body = {}

    chapter_id = str(body.get('chapterId') or '').strip()
    question_id = str(body.get('questionId') or '').strip()
    if not chapter_id or not question_id:
        return JsonResponse({'error': 'invalid payload'}, status=400)

    try:
        star_count = int(body.get('starCount') or 0)
    except (TypeError, ValueError):
        star_count = 0
    try:
        total_q = int(body.get('totalQuestions') or 0)
    except (TypeError, ValueError):
        total_q = 0

    prog, _ = LangChainProgress.objects.get_or_create(user=request.user)
    chapters = prog.chapters or {}
    ch = chapters.get(chapter_id) or {'completed': False, 'stars': 0, 'questions': {}}
    questions = ch.get('questions') or {}

    # 同一题只计一次分，重复提交不加星
    if questions.get(question_id):
        return JsonResponse({'alreadyDone': True, 'data': _langchain_progress_data(request.user, prog)})

    questions[question_id] = True
    ch['questions'] = questions
    ch['stars'] = int(ch.get('stars', 0)) + star_count
    prog.stars += star_count

    chapter_completed = False
    if not ch.get('completed'):
        denom = total_q or len(questions) or 1
        if len(questions) / denom >= 0.8:
            ch['completed'] = True
            prog.stars += 5  # 章节通关奖励，与前端本地逻辑保持一致
            chapter_completed = True

    chapters[chapter_id] = ch
    prog.chapters = chapters
    prog.save(update_fields=['stars', 'chapters', 'updated_at'])

    new_stars = star_count + (5 if chapter_completed else 0)
    payload = _langchain_progress_data(request.user, prog)
    return JsonResponse({
        'alreadyDone': False,
        'newStars': new_stars,
        'chapterCompleted': chapter_completed,
        'data': payload,
    })


@csrf_exempt
@_langchain_api_required
def langchain_api_spend(request):
    """消耗星星（答错换题等），星星不足时返回 success=false。"""
    if request.method != 'POST':
        return JsonResponse({'error': 'method not allowed'}, status=405)
    try:
        body = json.loads(request.body or b'{}')
    except json.JSONDecodeError:
        body = {}
    try:
        count = int(body.get('count') or 1)
    except (TypeError, ValueError):
        count = 1
    count = max(0, count)

    prog, _ = LangChainProgress.objects.get_or_create(user=request.user)
    data = _langchain_progress_data(request.user, prog)
    if prog.stars < count:
        return JsonResponse({'success': False, 'reason': 'stars', 'data': data})

    prog.stars -= count
    prog.save(update_fields=['stars', 'updated_at'])
    data = _langchain_progress_data(request.user, prog)
    return JsonResponse({'success': True, 'data': data})
